"""
Unit tests for Milestone 1: Core Extractors, Bundler & Tokenizer Remediation
===========================================================================
Validates:
- Patent classification regex false positives vs genuine patents
- Image extension recognition (.tif, .tiff, etc.)
- Multi-word companion SI suffix stripping in compute_bundle_id
- DOCX headings inclusion in sentences
- Sentence tokenization abbreviation handling (terminal vs mid-sentence)
- PDF document resource safety (clean handle closure on Windows)
- PPTX slides and table extraction
- Filename metadata parsing and underscore normalization
- DOI and Abstract extractor input guards and extraction
- Thesaurus logging, graceful WordNet fallback, and edge cases
"""

import logging
import os
import tempfile
from pathlib import Path
from unittest.mock import MagicMock, patch

import docx
import pptx
import pytest
from reportlab.pdfgen import canvas

from lws_engine.bundler import (
    classify_document_role,
    compute_bundle_id,
    bundle_document_record,
)
from lws_engine.extractors import (
    split_into_sentences,
    extract_pdf,
    extract_docx,
    extract_pptx,
    parse_filename_metadata,
    extract_doi,
    extract_abstract,
)
from lws_engine.thesaurus import ThesaurusManager


# ---------------------------------------------------------------------------
# 1. Bundler Classification & False Positive Tests
# ---------------------------------------------------------------------------

def test_bundler_patent_false_positives():
    """Verify standard academic filenames with 'us' + year do NOT classify as PATENT."""
    false_positive_candidates = [
        "focus2020.pdf",
        "status2021.pdf",
        "step2020.pdf",
        "two2020.pdf",
        "campus2020.pdf",
        "circus2022.pdf",
        "census2023.pdf",
        "focus_2020_optics.pdf",
        "status_2021_report.pdf",
    ]
    for filename in false_positive_candidates:
        role = classify_document_role(filename)
        assert role == "MAIN", f"Expected 'MAIN' for {filename}, got '{role}'"


def test_bundler_patent_true_positives():
    """Verify authentic patent identifiers are classified as PATENT."""
    patent_candidates = [
        "US20210053284A1.pdf",
        "US10123456.pdf",
        "WO2020123456A1.pdf",
        "EP2012345B1.pdf",
        "US-2021-0053284.pdf",
        "2021_US20210053284A1.pdf",
        "Additive_Manufacturing_Patent.pdf",
        "US11000123.pdf",
        "WO2019012345.pdf",
    ]
    for filename in patent_candidates:
        role = classify_document_role(filename)
        assert role == "PATENT", f"Expected 'PATENT' for {filename}, got '{role}'"


def test_bundler_image_extensions():
    """Verify all supported image formats including .tif and .tiff classify as IMAGE."""
    image_files = [
        "microscopy.tif",
        "diagram.tiff",
        "sample.png",
        "photo.jpg",
        "render.jpeg",
        "graph.bmp",
        "preview.webp",
    ]
    for filename in image_files:
        role = classify_document_role(filename)
        assert role == "IMAGE", f"Expected 'IMAGE' for {filename}, got '{role}'"


def test_bundler_peer_review_and_slides():
    """Verify peer review and presentation classifications."""
    assert classify_document_role("paper_reviewer_comments.pdf") == "PEER_REVIEW"
    assert classify_document_role("paper_peer_review.pdf") == "PEER_REVIEW"
    assert classify_document_role("editorial_decision_letter.pdf") == "PEER_REVIEW"
    assert classify_document_role("conference_talk.pptx") == "SLIDES"
    assert classify_document_role("seminar.ppt") == "SLIDES"


def test_compute_bundle_id_multi_word_si_suffixes():
    """Verify multi-word companion SI suffixes are stripped and match main paper bundle ID."""
    folder = "Tomography/2021"
    main_filename = "Kelly2021_Tomographic_Printing.pdf"
    main_id = compute_bundle_id(f"{folder}/{main_filename}", main_filename)
    assert main_id == "Tomography/2021/kelly2021-tomographic-printing"

    companion_variants = [
        "Kelly2021_Tomographic_Printing_supplementary_information.pdf",
        "Kelly2021_Tomographic_Printing_supplementary_info.pdf",
        "Kelly2021_Tomographic_Printing_supp_info.pdf",
        "Kelly2021_Tomographic_Printing_supporting_information.pdf",
        "Kelly2021_Tomographic_Printing_supporting_info.pdf",
        "Kelly2021_Tomographic_Printing_supplement.pdf",
        "Kelly2021_Tomographic_Printing-si.pdf",
        "Kelly2021_Tomographic_Printing (si).pdf",
        "Kelly2021_Tomographic_Printing [supplementary information].pdf",
        "Kelly2021_Tomographic_Printing_si_v2.pdf",
        "Kelly2021_Tomographic_Printing_peer_review.pdf",
        "Kelly2021_Tomographic_Printing_reviewer_comments.pdf",
    ]

    for comp in companion_variants:
        comp_id = compute_bundle_id(f"{folder}/{comp}", comp)
        assert comp_id == main_id, f"Bundle ID mismatch for {comp}: got '{comp_id}', expected '{main_id}'"


def test_bundle_document_record_helper():
    """Verify bundle_document_record returns role and bundle_id dictionary."""
    rec = bundle_document_record("papers/nature2020.pdf", "nature2020.pdf")
    assert rec["role"] == "MAIN"
    assert rec["bundle_id"] == "papers/nature2020"

    rec_si = bundle_document_record("papers/nature2020_supplementary_info.pdf", "nature2020_supplementary_info.pdf")
    assert rec_si["role"] == "SI"
    assert rec_si["bundle_id"] == "papers/nature2020"


# ---------------------------------------------------------------------------
# 2. Sentence Tokenization & Abbreviation Tests
# ---------------------------------------------------------------------------

def test_split_into_sentences_terminal_et_al():
    """Verify sentence splitting when 'et al.' terminates a sentence followed by a capital letter."""
    text = (
        "Volumetric 3D printing was pioneered by Loterie et al. "
        "We expanded upon their optical setup using a 405 nm laser. "
        "Further improvements were described by Kelly et al. "
        "High resolution resin formulations were developed."
    )
    sentences = split_into_sentences(text)
    assert len(sentences) == 4
    assert sentences[0] == "Volumetric 3D printing was pioneered by Loterie et al."
    assert sentences[1] == "We expanded upon their optical setup using a 405 nm laser."
    assert sentences[2] == "Further improvements were described by Kelly et al."
    assert sentences[3] == "High resolution resin formulations were developed."


def test_split_into_sentences_mid_sentence_citations():
    """Verify 'et al.' in the middle of a sentence does not trigger improper split."""
    text = "As Loterie et al. demonstrated in their 2020 work, resin cures rapidly."
    sentences = split_into_sentences(text)
    assert len(sentences) == 1
    assert sentences[0] == "As Loterie et al. demonstrated in their 2020 work, resin cures rapidly."


def test_split_into_sentences_quotes_and_brackets():
    """Verify sentence boundaries followed by quotation marks or bracketed numbers."""
    text = (
        'The primary finding was reported by Loterie et al. '
        '"We observed no optical aberrations," the authors noted.'
    )
    sentences = split_into_sentences(text)
    assert len(sentences) == 2
    assert sentences[0] == "The primary finding was reported by Loterie et al."
    assert '"We observed no optical aberrations," the authors noted.' in sentences[1]


def test_split_into_sentences_scientific_abbreviations():
    """Verify standard scientific abbreviations are protected."""
    text = (
        "Refer to Fig. 1 and Figs. 2 for calibration data. "
        "The reaction rate is given by Eq. 3 and Eq. 4. "
        "See Ref. 12 and Refs. 14 for background on Dr. Smith and Prof. Miller. "
        "Concentration was ca. 0.5 wt% vs. 1.0 wt% in vol. 5, no. 2, pp. 101-110."
    )
    sentences = split_into_sentences(text)
    assert len(sentences) == 4
    assert "Fig. 1 and Figs. 2" in sentences[0]
    assert "Eq. 3 and Eq. 4" in sentences[1]
    assert "Dr. Smith and Prof. Miller" in sentences[2]
    assert "ca. 0.5 wt% vs. 1.0 wt%" in sentences[3]


def test_split_into_sentences_edge_cases():
    """Verify input guards on None, non-string, empty, and short whitespace strings."""
    assert split_into_sentences(None) == []
    assert split_into_sentences("") == []
    assert split_into_sentences("   \n\t  ") == []
    assert split_into_sentences(12345) == []
    # Very short strings (<= 8 chars) are discarded as non-sentences
    assert split_into_sentences("Short.") == []


# ---------------------------------------------------------------------------
# 3. Document Extractors Tests (DOCX, PDF, PPTX)
# ---------------------------------------------------------------------------

def test_extract_docx_headings_included_in_sentences():
    """Verify DOCX Headings and Titles are captured in structured_pages sentences."""
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        docx_path = tmp.name

    try:
        doc = docx.Document()
        doc.add_heading("Tomographic Lithography in Photopolymers", level=0)
        doc.add_heading("1. Introduction", level=1)
        doc.add_paragraph("Volumetric manufacturing cures parts within seconds across all angles.")
        doc.add_heading("2. Experimental Methods", level=1)
        doc.add_heading("Optical Projection Setup", level=2)
        doc.add_paragraph("We used custom acrylates cured at 405 nm. This builds on Kelly et al. We achieved 50 um resolution.")
        doc.save(docx_path)

        full_text, structured_pages, meta = extract_docx(docx_path)

        assert meta["format"] == "DOCX"
        assert len(structured_pages) == 1
        sentences = structured_pages[0]["sentences"]

        # Headings must be present in sentences list
        assert "Tomographic Lithography in Photopolymers" in sentences
        assert "1. Introduction" in sentences
        assert "2. Experimental Methods" in sentences
        assert "Optical Projection Setup" in sentences

        # Paragraph sentences must also be present
        assert any("Volumetric manufacturing cures parts" in s for s in sentences)
        assert any("405 nm" in s for s in sentences)
        assert any("Kelly et al." in s for s in sentences)
        assert any("We achieved 50 um resolution." in s for s in sentences)
    finally:
        if os.path.exists(docx_path):
            os.remove(docx_path)


def test_extract_pdf_resource_safety():
    """Verify extract_pdf cleanly releases file handle using context management on Windows."""
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
        pdf_path = tmp.name

    try:
        # Create synthetic PDF
        c = canvas.Canvas(pdf_path)
        c.drawString(72, 750, "Computed Axial Lithography Document.")
        c.drawString(72, 730, "Loterie et al. demonstrated rapid printing. We validated their speed.")
        c.showPage()
        c.drawString(72, 750, "Second Page Heading")
        c.drawString(72, 730, "Additional results for photopolymerization kinetics.")
        c.showPage()
        c.save()

        full_text, structured_pages, meta = extract_pdf(pdf_path)

        assert meta["format"] == "PDF"
        assert meta["pages"] == 2
        assert len(structured_pages) == 2
        assert "Computed Axial Lithography" in structured_pages[0]["text"]
        assert len(structured_pages[0]["sentences"]) >= 2

        # Verify file can be removed immediately without Windows file-locking PermissionError
        os.remove(pdf_path)
        assert not os.path.exists(pdf_path)
    finally:
        if os.path.exists(pdf_path):
            os.remove(pdf_path)


def test_extract_pptx_slides_and_tables():
    """Verify PPTX slide titles, bullet points, and tables are extracted into sentences."""
    with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as tmp:
        pptx_path = tmp.name

    try:
        prs = pptx.Presentation()
        blank_slide_layout = prs.slide_layouts[6]

        # Slide 1: Title and bullet points
        slide1 = prs.slides.add_slide(prs.slide_layouts[0])
        slide1.shapes.title.text = "Volumetric 3D Printing Overview"
        subtitle = slide1.placeholders[1]
        subtitle.text = "Fast manufacturing of complex geometries."

        # Slide 2: Table
        slide2 = prs.slides.add_slide(blank_slide_layout)
        table_shape = slide2.shapes.add_table(2, 2, pptx.util.Inches(1), pptx.util.Inches(1), pptx.util.Inches(4), pptx.util.Inches(2))
        table = table_shape.table
        table.cell(0, 0).text = "Polymer"
        table.cell(0, 1).text = "Viscosity (mPa.s)"
        table.cell(1, 0).text = "PEGDA 250"
        table.cell(1, 1).text = "15.5"

        prs.save(pptx_path)

        full_text, structured_pages, meta = extract_pptx(pptx_path)

        assert meta["format"] == "PPTX"
        assert meta["pages"] == 2
        assert len(structured_pages) == 2

        # Check slide 1
        s1_sentences = structured_pages[0]["sentences"]
        assert any("Volumetric 3D Printing Overview" in s for s in s1_sentences)
        assert any("Fast manufacturing of complex geometries." in s for s in s1_sentences)

        # Check slide 2 table
        s2_sentences = structured_pages[1]["sentences"]
        assert any("PEGDA 250" in s for s in s2_sentences)
        assert any("Viscosity" in s for s in s2_sentences)
    finally:
        if os.path.exists(pptx_path):
            os.remove(pptx_path)


# ---------------------------------------------------------------------------
# 4. Filename Metadata & Academic Extractor Tests
# ---------------------------------------------------------------------------

def test_parse_filename_metadata_underscore_normalization():
    """Verify parse_filename_metadata normalizes underscores across all filename patterns."""
    cases = [
        (
            "2020 - Loterie - High_Resolution_VAM.pdf",
            {"year": 2020, "authors": "Loterie", "title": "High Resolution VAM"},
        ),
        (
            "2021 - Volumetric_Printing_Methods.docx",
            {"year": 2021, "authors": None, "title": "Volumetric Printing Methods"},
        ),
        (
            "Kelly - 2022 - Advanced_Tomography_Printing.pdf",
            {"year": 2022, "authors": "Kelly", "title": "Advanced Tomography Printing"},
        ),
        (
            "general_study_2019_results.pdf",
            {"year": 2019, "authors": None, "title": "general study 2019 results"},
        ),
    ]
    for filename, expected in cases:
        meta = parse_filename_metadata(filename)
        assert meta["year"] == expected["year"]
        assert meta["title"] == expected["title"]
        if expected["authors"] is not None:
            assert meta["authors"] == expected["authors"]

    # Test edge case inputs
    assert parse_filename_metadata(None) == {"year": None, "authors": None, "title": ""}
    assert parse_filename_metadata("") == {"year": None, "authors": None, "title": ""}
    assert parse_filename_metadata(12345) == {"year": None, "authors": None, "title": ""}


def test_extract_doi_guards_and_extraction():
    """Verify extract_doi handles valid DOIs, punctuation trimming, and input guards."""
    assert extract_doi(None) is None
    assert extract_doi("") is None
    assert extract_doi(12345) is None
    assert extract_doi("No DOI mentioned in this sentence.") is None

    # Valid DOIs with trailing punctuation
    sample1 = "Published at https://doi.org/10.1038/s41586-020-03079-9."
    assert extract_doi(sample1) == "10.1038/s41586-020-03079-9"

    sample2 = "DOI: 10.1126/science.abc1234; accessed June 2022"
    assert extract_doi(sample2) == "10.1126/science.abc1234"


def test_extract_abstract_guards_and_extraction():
    """Verify extract_abstract handles labeled abstracts and fallback chunks cleanly."""
    assert extract_abstract(None) is None
    assert extract_abstract("") is None
    assert extract_abstract(12345) is None

    labeled_text = (
        "Title of Paper\n\n"
        "Abstract — Volumetric additive manufacturing enables rapid fabrication of centimetre-scale objects "
        "within seconds via tomographic optical projection into photopolymer resins.\n\n"
        "1. Introduction\n"
        "Traditional 3D printing operates layer-by-layer."
    )
    abstract = extract_abstract(labeled_text)
    assert abstract is not None
    assert "Volumetric additive manufacturing enables rapid fabrication" in abstract
    assert "Traditional 3D printing" not in abstract

    # Fallback to leading introductory paragraph
    fallback_text = (
        "# Title\n\n"
        "Computed axial lithography uses multiple tomographic projections of light to solidify photocurable resin. "
        "This approach avoids oxygen inhibition layers and achieves smooth surface finishes without stair-stepping artifacts."
    )
    fallback_abstract = extract_abstract(fallback_text)
    assert fallback_abstract is not None
    assert "Computed axial lithography uses multiple tomographic projections" in fallback_abstract


# ---------------------------------------------------------------------------
# 5. Thesaurus Manager Tests (Logging, Fallback, Edge Cases)
# ---------------------------------------------------------------------------

def test_thesaurus_input_edge_cases():
    """Verify ThesaurusManager methods safely handle None, empty, and invalid max_count."""
    tm = ThesaurusManager(enable_wordnet=True)

    # Input guards for get_scientific_synonyms
    assert tm.get_scientific_synonyms(None) == []
    assert tm.get_scientific_synonyms("") == []
    assert tm.get_scientific_synonyms("   ") == []
    assert tm.get_scientific_synonyms(123) == []

    # Input guards for get_wordnet_synonyms
    assert tm.get_wordnet_synonyms(None) == []
    assert tm.get_wordnet_synonyms("") == []
    assert tm.get_wordnet_synonyms("resin", max_count=0) == []
    assert tm.get_wordnet_synonyms("resin", max_count=-5) == []
    assert tm.get_wordnet_synonyms(123) == []

    # Input guards for expand_query
    assert tm.expand_query(None) == {}
    assert tm.expand_query("") == {}
    assert tm.expand_query("   ") == {}
    assert tm.expand_query(123) == {}


def test_thesaurus_wordnet_logging_on_exception(caplog):
    """Verify ThesaurusManager logs a warning and returns gracefully when WordNet raises."""
    tm = ThesaurusManager(enable_wordnet=True)

    # Mock WordNet synsets to raise an exception
    mock_wn = MagicMock()
    mock_wn.synsets.side_effect = RuntimeError("Simulated WordNet disk read failure")
    tm._wordnet = mock_wn
    tm._wordnet_attempted = True

    with caplog.at_level(logging.WARNING, logger="lws_engine.thesaurus"):
        result = tm.get_wordnet_synonyms("polymer")

    assert result == []
    assert any("Failed to retrieve WordNet synonyms for 'polymer'" in record.message for record in caplog.records)


def test_thesaurus_wordnet_unavailable_logging(caplog):
    """Verify ThesaurusManager logs warning when WordNet corpus is unavailable."""
    tm = ThesaurusManager(enable_wordnet=True)
    tm._wordnet_attempted = False

    # Force _get_wordnet to trigger exception
    with patch("nltk.corpus.wordnet.synsets", side_effect=LookupError("WordNet corpus missing")):
        with caplog.at_level(logging.WARNING, logger="lws_engine.thesaurus"):
            wn = tm._get_wordnet()
            assert wn is None

    assert any("WordNet corpus is unavailable" in record.message for record in caplog.records)


def test_thesaurus_bidirectional_and_multiword():
    """Verify bidirectional scientific synonyms and multi-word phrase query expansion."""
    tm = ThesaurusManager(enable_wordnet=False)

    # Reverse lookup: 'silicone' should contain 'pdms'
    silicone_syns = tm.get_scientific_synonyms("silicone")
    assert "pdms" in silicone_syns or "polydimethylsiloxane" in silicone_syns

    # Multi-word lookup: 'computed axial lithography' should match 'cal' / 'vam'
    cal_syns = tm.get_scientific_synonyms("computed axial lithography")
    assert any("vam" in s or "volumetric" in s for s in cal_syns)

    # Multi-word query expansion
    expanded = tm.expand_query("refractive index matching", use_wordnet=False, use_scientific=True)
    assert any("refractive" in key for key in expanded.keys())
