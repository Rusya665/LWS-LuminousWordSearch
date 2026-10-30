"""
LWS 2.0 Document & Metadata Extractors
High-speed extraction for PDF (pypdfium2), DOCX, PPTX, and text segmentation into sentences.
"""

from __future__ import annotations

import hashlib
import json
import re
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, Union

from lws_engine.config import win_safe_path

try:
    import pypdfium2 as pdfium
except ImportError:
    pdfium = None

try:
    import pptx
except ImportError:
    pptx = None


def split_into_sentences(text: str) -> List[str]:
    """
    Splits text into clean sentences without breaking on common abbreviations:
    (e.g., 'et al.', 'Fig.', 'i.e.', 'Dr.', 'Eq.', 'vs.').
    Preserves sentence boundaries when an abbreviation (e.g. 'et al.') is followed
    by a capitalized sentence start.
    """
    if not text or not isinstance(text, str):
        return []

    # Normalize whitespace
    cleaned = re.sub(r"[\r\n]+", " ", text).strip()
    if not cleaned:
        return []
    cleaned = re.sub(r"\s+", " ", cleaned)

    # Protect abbreviations: et al. only protected when NOT followed by capitalized sentence start
    placeholders = {
        r"\bet al\.(?-i:(?!\s+[\"'\u201c\u2018]?[A-Z]))": "ET_AL_TOKEN",
        r"\be\.g\.": "EG_TOKEN",
        r"\bi\.e\.": "IE_TOKEN",
        r"\bFig\.(?=\s*[\(\[]?(?:S\d+|\d+(?:\.\d+)?[a-zA-Z]?|[IVXLCDM]+))": "FIG_TOKEN",
        r"\bFigs\.(?=\s*[\(\[]?(?:S\d+|\d+(?:\.\d+)?[a-zA-Z]?|[IVXLCDM]+))": "FIGS_TOKEN",
        r"\bEq\.(?=\s*[\(\[]?(?:S\d+|\d+(?:\.\d+)?[a-zA-Z]?|[IVXLCDM]+))": "EQ_TOKEN",
        r"\bEqs\.(?=\s*[\(\[]?(?:S\d+|\d+(?:\.\d+)?[a-zA-Z]?|[IVXLCDM]+))": "EQS_TOKEN",
        r"\bRef\.(?=\s*[\(\[]?(?:S\d+|\d+(?:\.\d+)?[a-zA-Z]?|[IVXLCDM]+))": "REF_TOKEN",
        r"\bRefs\.(?=\s*[\(\[]?(?:S\d+|\d+(?:\.\d+)?[a-zA-Z]?|[IVXLCDM]+))": "REFS_TOKEN",
        r"\bDr\.": "DR_TOKEN",
        r"\bProf\.": "PROF_TOKEN",
        r"\bMr\.": "MR_TOKEN",
        r"\bMrs\.": "MRS_TOKEN",
        r"\bMs\.": "MS_TOKEN",
        r"\bvol\.": "VOL_TOKEN",
        r"\bno\.\s*(?=\d+)": "NO_TOKEN ",
        r"\bpp\.": "PP_TOKEN",
        r"\bvs\.": "VS_TOKEN",
        r"\bca\.": "CA_TOKEN",
    }
    protected = cleaned
    for pattern, repl in placeholders.items():
        protected = re.sub(pattern, repl, protected, flags=re.IGNORECASE)

    # Split on sentence boundaries (. ! ?) allowing optional closing quotes/brackets/parentheses
    split_pattern = r"(?:(?<=[.!?])|(?<=[.!?][\"'\u201d\u2019\)\],]))\s+(?=[A-Z0-9\(\[\"'\u201c\u2018\u2022\u25aa\u2192\-\*\U00010000-\U0010ffff])"
    raw_sents = re.split(split_pattern, protected)

    sentences: List[str] = []
    for s in raw_sents:
        # Restore placeholders
        restored = s.replace("ET_AL_TOKEN", "et al.")
        restored = restored.replace("EG_TOKEN", "e.g.")
        restored = restored.replace("IE_TOKEN", "i.e.")
        restored = restored.replace("FIGS_TOKEN", "Figs.")
        restored = restored.replace("FIG_TOKEN", "Fig.")
        restored = restored.replace("EQS_TOKEN", "Eqs.")
        restored = restored.replace("EQ_TOKEN", "Eq.")
        restored = restored.replace("REFS_TOKEN", "Refs.")
        restored = restored.replace("REF_TOKEN", "Ref.")
        restored = restored.replace("DR_TOKEN", "Dr.")
        restored = restored.replace("PROF_TOKEN", "Prof.")
        restored = restored.replace("MR_TOKEN", "Mr.")
        restored = restored.replace("MRS_TOKEN", "Mrs.")
        restored = restored.replace("MS_TOKEN", "Ms.")
        restored = restored.replace("VOL_TOKEN", "vol.")
        restored = restored.replace("NO_TOKEN ", "no. ")
        restored = restored.replace("PP_TOKEN", "pp.")
        restored = restored.replace("VS_TOKEN", "vs.")
        restored = restored.replace("CA_TOKEN", "ca.")

        trimmed = restored.strip()
        if len(trimmed) > 8:
            sentences.append(trimmed)

    return sentences


def extract_pdf(file_path: Union[str, Path]) -> Tuple[str, List[Dict[str, Any]], Dict[str, Any]]:
    """
    Extracts text and page-by-page sentence breakdown using pypdfium2.
    Returns: (full_text_markdown, structured_pages, format_meta)
    """
    if pdfium is None:
        raise RuntimeError("pypdfium2 is required for fast PDF extraction.")

    safe_path = win_safe_path(file_path)
    pages_output: List[str] = []
    structured_pages: List[Dict[str, Any]] = []

    with pdfium.PdfDocument(str(safe_path)) as doc:
        num_pages = len(doc)
        for page_idx in range(num_pages):
            page = doc[page_idx]
            textpage = page.get_textpage()
            page_text = textpage.get_text_range().strip()

            # Heading inference from page header
            first_line = page_text.splitlines()[0].strip() if page_text else ""
            heading = first_line[:100] if len(first_line) > 3 else f"Page {page_idx + 1}"

            sentences = split_into_sentences(page_text)
            structured_pages.append({
                "page_num": page_idx + 1,
                "heading": heading,
                "text": page_text,
                "sentences": sentences,
            })

            if page_text:
                pages_output.append(f"## Page {page_idx + 1}\n\n{page_text}")

    full_text = "\n\n".join(pages_output)
    meta = {
        "format": "PDF",
        "pages": num_pages,
        "char_count": len(full_text),
    }
    return full_text, structured_pages, meta


def extract_docx(file_path: Union[str, Path]) -> Tuple[str, List[Dict[str, Any]], Dict[str, Any]]:
    """Extracts text and headings from a DOCX file without heavy dependencies."""
    paragraphs: List[str] = []
    sentences: List[str] = []
    current_heading = "Document"
    safe_fp = win_safe_path(file_path)

    with zipfile.ZipFile(safe_fp) as z:
        xml_content = z.read("word/document.xml")

    tree = ET.fromstring(xml_content)
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    val_key = f"{{{ns['w']}}}val"

    def _extract_table_rows(tbl_elem) -> List[str]:
        table_lines = []
        for tr in tbl_elem.findall("w:tr", ns):
            row_cells = []
            for tc in tr.findall("w:tc", ns):
                for inner_tbl in tc.findall("w:tbl", ns):
                    table_lines.extend(_extract_table_rows(inner_tbl))
                cell_parts = []
                for p in tc.findall("w:p", ns):
                    texts = [node.text for node in p.iterfind(".//w:t", ns) if node.text]
                    t = "".join(texts).strip()
                    if t:
                        cell_parts.append(t)
                if cell_parts:
                    row_cells.append(" ".join(cell_parts))
            if row_cells:
                table_lines.append(" | ".join(row_cells))
        return table_lines

    def _process_container(container) -> List[Tuple[str, str]]:
        elements: List[Tuple[str, str]] = []
        for child in container:
            tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
            if tag == "p":
                p_style = child.find(".//w:pPr/w:pStyle", ns)
                style_val = p_style.attrib.get(val_key, "") if p_style is not None else ""
                s_val_lower = style_val.lower()

                texts = [node.text for node in child.iterfind(".//w:t", ns) if node.text]
                if texts:
                    line = "".join(texts).strip()
                    if line:
                        if any(h in s_val_lower for h in ["heading1", "heading 1", "title"]):
                            elements.append(("h1", line))
                        elif "heading" in s_val_lower:
                            elements.append(("h", line))
                        else:
                            elements.append(("p", line))
            elif tag == "tbl":
                tbl_rows = _extract_table_rows(child)
                for tr_line in tbl_rows:
                    elements.append(("tbl", tr_line))
            elif tag == "sdt":
                sdt_content = child.find("w:sdtContent", ns)
                if sdt_content is not None:
                    elements.extend(_process_container(sdt_content))
        return elements

    body = tree.find("w:body", ns)
    if body is None:
        body = tree

    for kind, line in _process_container(body):
        if kind == "h1":
            current_heading = line
            paragraphs.append(f"## {line}")
            sents = split_into_sentences(line)
            if len(sents) > 1:
                sentences.extend(sents)
            else:
                sentences.append(line)
        elif kind == "h":
            paragraphs.append(f"### {line}")
            sents = split_into_sentences(line)
            if len(sents) > 1:
                sentences.extend(sents)
            else:
                sentences.append(line)
        elif kind == "tbl":
            paragraphs.append(line)
            sents = split_into_sentences(line)
            if sents:
                sentences.extend(sents)
            elif len(line) >= 3:
                sentences.append(line)
        else:
            paragraphs.append(line)
            sentences.extend(split_into_sentences(line))

    full_text = "\n\n".join(paragraphs)
    structured_pages = [{
        "page_num": 1,
        "heading": current_heading,
        "text": full_text,
        "sentences": sentences,
    }]
    meta = {
        "format": "DOCX",
        "paragraphs": len(paragraphs),
        "char_count": len(full_text),
        "pages": 1,
    }
    return full_text, structured_pages, meta


def extract_pptx(file_path: Union[str, Path]) -> Tuple[str, List[Dict[str, Any]], Dict[str, Any]]:
    """Extracts text from slides in a PPTX file using python-pptx."""
    if pptx is None:
        raise RuntimeError("python-pptx is required for PPTX extraction.")

    safe_fp = win_safe_path(file_path)
    prs = pptx.Presentation(str(safe_fp))
    slides_text: List[str] = []
    structured_pages: List[Dict[str, Any]] = []

    for i, slide in enumerate(prs.slides):
        slide_lines: List[str] = []
        slide_title = f"Slide {i + 1}"
        if slide.shapes.title and slide.shapes.title.text.strip():
            slide_title = slide.shapes.title.text.strip()

        def _extract_shape_lines(shape) -> List[str]:
            lines: List[str] = []
            if getattr(shape, "has_text_frame", False):
                try:
                    for paragraph in shape.text_frame.paragraphs:
                        line = paragraph.text.strip()
                        if line:
                            lines.append(line)
                except Exception:
                    pass
            if getattr(shape, "has_table", False):
                try:
                    for row in shape.table.rows:
                        row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_texts:
                            lines.append(" | ".join(row_texts))
                except Exception:
                    pass
            if hasattr(shape, "shapes"):
                try:
                    for child in shape.shapes:
                        lines.extend(_extract_shape_lines(child))
                except Exception:
                    pass
            return lines

        for shape in slide.shapes:
            slide_lines.extend(_extract_shape_lines(shape))

        slide_body = "\n".join(f"- {l}" if not l.startswith("-") else l for l in slide_lines)
        sentences: List[str] = []
        for line in slide_lines:
            sents = split_into_sentences(line)
            if sents:
                sentences.extend(sents)
            elif len(line) >= 3:
                sentences.append(line)

        structured_pages.append({
            "page_num": i + 1,
            "heading": slide_title,
            "text": slide_body,
            "sentences": sentences,
        })
        if slide_lines:
            slides_text.append(f"## {slide_title}\n\n{slide_body}")

    full_text = "\n\n".join(slides_text)
    meta = {
        "format": "PPTX",
        "pages": len(prs.slides),
        "char_count": len(full_text),
    }
    return full_text, structured_pages, meta


def parse_filename_metadata(filename: str) -> Dict[str, Any]:
    """Infers Year, Authors, and Title from academic filename conventions."""
    if not filename or not isinstance(filename, str):
        return {"year": None, "authors": None, "title": ""}

    stem = Path(filename).stem.strip()
    meta: Dict[str, Any] = {"year": None, "authors": None, "title": stem}

    m = re.match(r"^(\d{4})\s*-\s*([^-]+?)\s*-\s*(.+)$", stem)
    if m:
        meta["year"] = int(m.group(1))
        meta["authors"] = m.group(2).strip()
        meta["title"] = m.group(3).strip()
    else:
        m = re.match(r"^(\d{4})\s*-\s*(.+)$", stem)
        if m:
            meta["year"] = int(m.group(1))
            meta["title"] = m.group(2).strip()
        else:
            m = re.match(r"^([^-]+?)\s*-\s*(\d{4})\s*-\s*(.+)$", stem)
            if m:
                meta["authors"] = m.group(1).strip()
                meta["year"] = int(m.group(2))
                meta["title"] = m.group(3).strip()
            else:
                year_match = re.search(r"(?<!\d)(19\d\d|20\d\d)(?!\d)", stem)
                if year_match:
                    meta["year"] = int(year_match.group(1))

    meta["title"] = meta["title"].replace("_", " ").strip()
    return meta


def extract_doi(text: str) -> Optional[str]:
    """Finds the first valid Crossref/academic DOI in text."""
    if not text or not isinstance(text, str):
        return None
    match = re.search(r"\b10\.\d{4,9}/[-._;()/:A-Za-z0-9]+\b", text)
    if match:
        return match.group(0).rstrip(".,;")
    return None


def extract_abstract(text: str) -> Optional[str]:
    """Isolates the paper's Abstract or leading introductory paragraph."""
    if not text or not isinstance(text, str):
        return None
    match = re.search(
        r"(?im)^[ \t]*(?:#+\s*|\*+\s*)?abstract\b[ \t*]*(?:[:.—–-][ \t*]*)?\s*([\s\S]{40,2500}?)(?=\n\s*(?:(?:(?:1|\bI\b)\.?\s*)?introduction|keywords|index terms|background|\n\s*\n\s*\n|$))",
        text,
    )
    if match:
        return re.sub(r"\s+", " ", match.group(1).strip())

    cleaned = re.sub(r"^#+.*$", "", text, flags=re.MULTILINE).strip()
    for chunk in cleaned.split("\n\n"):
        chunk = chunk.strip()
        if len(chunk) > 60 and not chunk.startswith("<!--"):
            snippet = re.sub(r"\s+", " ", chunk)[:450]
            return snippet + "..." if len(chunk) > 450 else snippet
    return None
